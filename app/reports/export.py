from __future__ import annotations

import io
import zipfile
from typing import Dict, Any
import pandas as pd
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList


def dataframes_to_csv_bytes(results: Dict[str, Any]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, mode="w", compression=zipfile.ZIP_DEFLATED) as zf:
        for name in [
            ("realized_lots.csv", results["realized_lots"]),
            ("per_scrip_summary.csv", results["per_scrip_summary"]),
            ("overall_summary.csv", results["overall_summary"]),
            ("open_positions.csv", results["open_positions"]),
        ]:
            zf.writestr(name[0], name[1].to_csv(index=False))
    return buf.getvalue()


def dataframes_to_excel_bytes(results: Dict[str, Any]) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        # Write all sheets
        results["realized_lots"].to_excel(writer, index=False, sheet_name="RealizedLots")
        results["per_scrip_summary"].to_excel(writer, index=False, sheet_name="PerScripSummary")
        results["overall_summary"].to_excel(writer, index=False, sheet_name="OverallSummary")
        results["open_positions"].to_excel(writer, index=False, sheet_name="OpenPositions")

        # Access workbook for chart creation
        workbook = writer.book

        # Add charts to PerScripSummary sheet
        _add_per_scrip_charts(workbook, results["per_scrip_summary"])

        # Add charts to OverallSummary sheet
        _add_overall_charts(workbook, results["overall_summary"])

    return buf.getvalue()


def _add_per_scrip_charts(workbook, per_scrip_df: pd.DataFrame):
    """Add STCG/LTCG bar chart to PerScripSummary sheet"""
    if per_scrip_df.empty:
        return

    ws = workbook["PerScripSummary"]

    # Only create charts if we have data
    num_rows = len(per_scrip_df)
    if num_rows == 0:
        return

    # Bar Chart: STCG vs LTCG per Scrip
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "STCG vs LTCG by Scrip"
    chart.y_axis.title = "Gain (INR)"
    chart.x_axis.title = "Scrip"

    # Data references (assumes Scrip in col A, STCG_Total in col B, LTCG_Total in col C)
    data = Reference(ws, min_col=2, min_row=1, max_row=num_rows + 1, max_col=3)
    cats = Reference(ws, min_col=1, min_row=2, max_row=num_rows + 1)

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4

    # Position chart below the data
    ws.add_chart(chart, f"A{num_rows + 4}")

    # Net Gain Bar Chart
    net_chart = BarChart()
    net_chart.type = "col"
    net_chart.style = 11
    net_chart.title = "Net Total Gain by Scrip"
    net_chart.y_axis.title = "Net Gain (INR)"
    net_chart.x_axis.title = "Scrip"

    # Net_Total_Gain is in column D
    net_data = Reference(ws, min_col=4, min_row=1, max_row=num_rows + 1)

    net_chart.add_data(net_data, titles_from_data=True)
    net_chart.set_categories(cats)

    # Position chart to the right of first chart
    ws.add_chart(net_chart, f"J{num_rows + 4}")


def _add_overall_charts(workbook, overall_df: pd.DataFrame):
    """Add pie chart for overall STCG/LTCG distribution"""
    if overall_df.empty:
        return

    ws = workbook["OverallSummary"]

    # Get STCG and LTCG totals
    stcg_total = float(overall_df.iloc[0]["STCG_Total"])
    ltcg_total = float(overall_df.iloc[0]["LTCG_Total"])

    # Only create pie chart if there are gains to show
    if stcg_total > 0 or ltcg_total > 0:
        # Pie Chart: STCG vs LTCG distribution
        pie = PieChart()
        pie.title = "STCG vs LTCG Distribution"

        # Data references (STCG_Total in col A, LTCG_Total in col B)
        labels = Reference(ws, min_col=1, min_row=1, max_col=2)
        data = Reference(ws, min_col=1, min_row=2, max_col=2)

        pie.add_data(data, titles_from_data=False)
        pie.set_categories(labels)

        # Show data labels
        pie.dataLabels = DataLabelList()
        pie.dataLabels.showCatName = True
        pie.dataLabels.showVal = True

        # Position chart below the data
        ws.add_chart(pie, "A4")

